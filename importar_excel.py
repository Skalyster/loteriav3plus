#!/usr/bin/env python3
"""
Importa los históricos de Lotoideas (Excel) a loteria_db.sqlite con el formato de app+p+e.py.

  La Primitiva : numeros = "n1,...,n6"        reintegro = R (0 si no existe en el histórico)
  Euromillones : numeros = "n1,...,n5,e1,e2"  reintegro = 0

Uso:
    python importar_excel.py --db loteria_db.sqlite \
        --primitiva "Histórico de Resultados - Primitiva.xlsx" \
        --euromillones "Histórico de Resultados - Euromillones.xlsx"

Los sorteos que ya existen (mismo juego y fecha) se omiten, así que se puede repetir.
"""
import argparse
import sqlite3
from datetime import datetime

import openpyxl


def entero(v):
    try:
        return int(float(str(v).strip()))
    except (TypeError, ValueError):
        return None


def fecha_iso(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    return None


def leer_primitiva(ruta):
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    out, malas = [], 0
    for ws in wb.worksheets:
        for fila in ws.iter_rows(values_only=True):
            f = fecha_iso(fila[0] if fila else None)
            if not f:
                continue
            nums = [entero(x) for x in fila[1:7]]
            if None in nums or len(set(nums)) != 6 or not all(1 <= n <= 49 for n in nums):
                malas += 1
                continue
            r = entero(fila[8]) if len(fila) > 8 else None
            if r is None or not 0 <= r <= 9:
                r = 0
            out.append((f, ",".join(map(str, sorted(nums))), r))
    return out, malas


def leer_euromillones(ruta):
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    out, malas = [], 0
    for ws in wb.worksheets:
        for fila in ws.iter_rows(values_only=True):
            f = fecha_iso(fila[0] if fila else None)
            if not f:
                continue
            nums = [entero(x) for x in fila[1:6]]
            est = [entero(x) for x in fila[7:9]] if len(fila) > 8 else [None]
            if (None in nums or None in est or len(set(nums)) != 5 or len(set(est)) != 2
                    or not all(1 <= n <= 50 for n in nums) or not all(1 <= e <= 12 for e in est)):
                malas += 1
                continue
            out.append((f, ",".join(map(str, sorted(nums) + sorted(est))), 0))
    return out, malas


def volcar(cur, juego, filas):
    nuevos = omitidos = 0
    vistos = set()
    for f, nums, r in filas:
        if f in vistos:
            omitidos += 1
            continue
        vistos.add(f)
        if cur.execute("SELECT 1 FROM sorteos WHERE juego=? AND fecha=?", (juego, f)).fetchone():
            omitidos += 1
            continue
        cur.execute("INSERT INTO sorteos (juego, fecha, numeros, reintegro) VALUES (?,?,?,?)", (juego, f, nums, r))
        nuevos += 1
    return nuevos, omitidos


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--db", default="loteria_db.sqlite")
    ap.add_argument("--primitiva")
    ap.add_argument("--euromillones")
    a = ap.parse_args()

    con = sqlite3.connect(a.db)
    cur = con.cursor()
    if a.primitiva:
        filas, malas = leer_primitiva(a.primitiva)
        n, o = volcar(cur, "La Primitiva", filas)
        print(f"La Primitiva: {n} nuevos, {o} omitidos, {malas} filas ilegibles")
    if a.euromillones:
        filas, malas = leer_euromillones(a.euromillones)
        n, o = volcar(cur, "Euromillones", filas)
        print(f"Euromillones: {n} nuevos, {o} omitidos, {malas} filas ilegibles")
    con.commit()
    con.close()


if __name__ == "__main__":
    main()
